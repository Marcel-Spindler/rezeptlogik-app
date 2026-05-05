#!/usr/bin/env python
# -*- coding: utf-8 -*-

"""
KW-Tool: Transaction Log -> Excel Upsert + optional TSV-Export fuer Google Sheets.

Funktionen:
- Liest alle Wochen aus dem Transaction Log.
- Aggregiert Kennzahlen je KW, Produkt und Bereich.
- Schreibt in eine bestehende Excel-Datei (oder erstellt sie neu), ohne Duplikate.
- Aktualisiert bestehende Zeilen anhand eindeutiger Schluessel.
- Kann TSV-Dateien fuer Google Sheets erzeugen.
"""

from __future__ import annotations

import argparse
import csv
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Tuple

from openpyxl import Workbook, load_workbook


SUMMARY_SHEET = "KW_Summary"
PRODUCT_SHEET = "KW_Produkte"
AREA_SHEET = "KW_Bereiche"

SUMMARY_HEADERS = [
    "Week",
    "Transaktionen",
    "Produkte_Anzahl",
    "Bereiche_Anzahl",
    "Zugaenge_kg",
    "Bewegungen_kg",
    "Differenz_kg",
    "Netto_kg",
]

PRODUCT_HEADERS = [
    "Week",
    "Produkt",
    "Zugaenge_kg",
    "Bewegungen_kg",
    "Differenz_kg",
    "Netto_kg",
    "Grammatur_bis_BlastChiller_g",
    "Transaktionen",
    "Status",
]

AREA_HEADERS = [
    "Week",
    "Bereich",
    "Zugaenge_kg",
    "Bewegungen_kg",
    "Differenz_kg",
    "Netto_kg",
    "Kategorie",
]


@dataclass
class ProductAgg:
    pos: float = 0.0
    neg: float = 0.0
    net: float = 0.0
    trans: int = 0
    blast_g: float = 0.0


@dataclass
class AreaAgg:
    pos: float = 0.0
    neg: float = 0.0
    net: float = 0.0


@dataclass
class WeekAgg:
    trans: int = 0
    pos: float = 0.0
    neg: float = 0.0
    net: float = 0.0
    products: set = None
    areas: set = None

    def __post_init__(self) -> None:
        if self.products is None:
            self.products = set()
        if self.areas is None:
            self.areas = set()


def to_float(value) -> float:
    if value is None or value == "":
        return 0.0
    try:
        return float(value)
    except Exception:
        return 0.0


def normalize_week(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    return text


def category_for_area(area: str) -> str:
    al = area.lower()
    prep_kw = ("preb", "prep", "plating", "post", "sleeving", "deboxwip", "platingwip")
    fulfil_kw = ("fab", "fac", "preb", "pack", "slip", "vf-", "plh-", "psh-", "vegg")
    lager_kw = ("lager", "storage", "ambient", "chilled", "frozen", "pick", "slot", "stgdr")

    if any(k in al for k in prep_kw):
        return "Prep/Plating"
    if any(k in al for k in fulfil_kw):
        return "Fulfillment"
    if any(k in al for k in lager_kw):
        return "Lager"
    return "Andere"


def status_for_net(net: float) -> str:
    if net > 0:
        return "EINGANG +"
    if net < 0:
        return "AUSGANG -"
    return "NEUTRAL"


def ensure_sheet(wb, name: str, headers: List[str]):
    ws = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
    if ws.max_row == 1 and ws.cell(row=1, column=1).value is None:
        for i, h in enumerate(headers, start=1):
            ws.cell(row=1, column=i, value=h)
    else:
        existing = [ws.cell(row=1, column=i).value for i in range(1, len(headers) + 1)]
        if existing != headers:
            for i, h in enumerate(headers, start=1):
                ws.cell(row=1, column=i, value=h)
    return ws


def build_index(ws, key_columns: List[str]) -> Dict[Tuple[str, ...], int]:
    header_map = {}
    for col in range(1, ws.max_column + 1):
        name = ws.cell(row=1, column=col).value
        if name:
            header_map[str(name)] = col

    index = {}
    for row in range(2, ws.max_row + 1):
        key_values = []
        for key_col in key_columns:
            c = header_map.get(key_col)
            if not c:
                key_values = []
                break
            key_values.append(str(ws.cell(row=row, column=c).value or "").strip())
        if key_values:
            index[tuple(key_values)] = row
    return index


def upsert_rows(ws, headers: List[str], key_columns: List[str], rows: Iterable[List]):
    header_pos = {h: i + 1 for i, h in enumerate(headers)}
    key_idx = [headers.index(k) for k in key_columns]
    idx = build_index(ws, key_columns)

    for row_values in rows:
        key = tuple(str(row_values[i]).strip() for i in key_idx)
        if key in idx:
            row_no = idx[key]
        else:
            row_no = ws.max_row + 1
            idx[key] = row_no

        for i, value in enumerate(row_values, start=1):
            ws.cell(row=row_no, column=i, value=value)


def aggregate(transaction_log: Path, blast_patterns: List[str]):
    wb = load_workbook(transaction_log)
    ws = wb.active

    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    col_map = {str(h): i + 1 for i, h in enumerate(headers) if h is not None}

    required = ["Week", "Item Desc", "Tran Qty", "Location Id"]
    missing = [c for c in required if c not in col_map]
    if missing:
        raise ValueError(f"Pflichtspalten fehlen im Transaction Log: {missing}")

    week_agg: Dict[str, WeekAgg] = defaultdict(WeekAgg)
    product_agg: Dict[Tuple[str, str], ProductAgg] = defaultdict(ProductAgg)
    area_agg: Dict[Tuple[str, str], AreaAgg] = defaultdict(AreaAgg)

    patterns = [p.lower().strip() for p in blast_patterns if p and p.strip()]

    for row in range(2, ws.max_row + 1):
        week = normalize_week(ws.cell(row=row, column=col_map["Week"]).value)
        if not week:
            continue

        product = ws.cell(row=row, column=col_map["Item Desc"]).value
        if product is None or str(product).strip() == "":
            continue
        product = str(product).strip()

        qty = to_float(ws.cell(row=row, column=col_map["Tran Qty"]).value)

        area_raw = ws.cell(row=row, column=col_map["Location Id"]).value
        area = str(area_raw).strip() if area_raw else "UNBEKANNT"

        w = week_agg[week]
        w.trans += 1
        w.products.add(product)
        w.areas.add(area)

        pk = (week, product)
        ak = (week, area)

        if qty >= 0:
            w.pos += qty
            product_agg[pk].pos += qty
            area_agg[ak].pos += qty
        else:
            abs_q = abs(qty)
            w.neg += abs_q
            product_agg[pk].neg += abs_q
            area_agg[ak].neg += abs_q

        w.net += qty
        product_agg[pk].net += qty
        area_agg[ak].net += qty
        product_agg[pk].trans += 1

        area_l = area.lower()
        if any(p in area_l for p in patterns):
            product_agg[pk].blast_g += abs(qty) * 1000.0

    return week_agg, product_agg, area_agg


def to_summary_rows(week_agg: Dict[str, WeekAgg]) -> List[List]:
    rows = []
    for week in sorted(week_agg.keys()):
        w = week_agg[week]
        rows.append([
            week,
            w.trans,
            len(w.products),
            len(w.areas),
            round(w.pos, 2),
            round(w.neg, 2),
            round(w.pos - w.neg, 2),
            round(w.net, 2),
        ])
    return rows


def to_product_rows(product_agg: Dict[Tuple[str, str], ProductAgg]) -> List[List]:
    rows = []
    sorted_keys = sorted(product_agg.keys(), key=lambda k: (k[0], -abs(product_agg[k].net), k[1]))
    for week, product in sorted_keys:
        p = product_agg[(week, product)]
        rows.append([
            week,
            product,
            round(p.pos, 2),
            round(p.neg, 2),
            round(p.pos - p.neg, 2),
            round(p.net, 2),
            int(round(p.blast_g, 0)),
            p.trans,
            status_for_net(p.net),
        ])
    return rows


def to_area_rows(area_agg: Dict[Tuple[str, str], AreaAgg]) -> List[List]:
    rows = []
    sorted_keys = sorted(area_agg.keys(), key=lambda k: (k[0], -abs(area_agg[k].net), k[1]))
    for week, area in sorted_keys:
        a = area_agg[(week, area)]
        rows.append([
            week,
            area,
            round(a.pos, 2),
            round(a.neg, 2),
            round(a.pos - a.neg, 2),
            round(a.net, 2),
            category_for_area(area),
        ])
    return rows


def export_sheet_to_tsv(ws, target: Path):
    with target.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f, delimiter="\t")
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column, values_only=True):
            writer.writerow(["" if v is None else v for v in row])


def main():
    parser = argparse.ArgumentParser(description="Transaction Log -> Excel Upsert Tool")
    parser.add_argument("--transaction-log", required=True, help="Pfad zur Transaction_Log.xlsx")
    parser.add_argument(
        "--excel",
        required=True,
        help="Ziel-Excel (bestehend oder neu), die als Master fuer Google Sheets dient",
    )
    parser.add_argument(
        "--blast-patterns",
        default="blast,chill,postb-01",
        help="Kommagetrennte Location-Muster fuer Blast-Chiller-Grammatur",
    )
    parser.add_argument(
        "--export-tsv-dir",
        default="",
        help="Optionales Verzeichnis fuer TSV-Export der 3 Sheets",
    )

    args = parser.parse_args()

    tx_path = Path(args.transaction_log)
    excel_path = Path(args.excel)
    blast_patterns = [p.strip() for p in args.blast_patterns.split(",") if p.strip()]

    if not tx_path.exists():
        raise FileNotFoundError(f"Transaction Log nicht gefunden: {tx_path}")

    week_agg, product_agg, area_agg = aggregate(tx_path, blast_patterns)

    if excel_path.exists():
        wb = load_workbook(excel_path)
    else:
        wb = Workbook()
        default = wb.active
        wb.remove(default)

    ws_summary = ensure_sheet(wb, SUMMARY_SHEET, SUMMARY_HEADERS)
    ws_products = ensure_sheet(wb, PRODUCT_SHEET, PRODUCT_HEADERS)
    ws_areas = ensure_sheet(wb, AREA_SHEET, AREA_HEADERS)

    upsert_rows(ws_summary, SUMMARY_HEADERS, ["Week"], to_summary_rows(week_agg))
    upsert_rows(ws_products, PRODUCT_HEADERS, ["Week", "Produkt"], to_product_rows(product_agg))
    upsert_rows(ws_areas, AREA_HEADERS, ["Week", "Bereich"], to_area_rows(area_agg))

    wb.save(excel_path)

    print("=" * 90)
    print("KW TOOL ERFOLGREICH AUSGEFUEHRT")
    print("=" * 90)
    print(f"Transaction Log: {tx_path}")
    print(f"Master Excel:    {excel_path}")
    print(f"Wochen erkannt:  {len(week_agg)}")
    print(f"Produkte (keys): {len(product_agg)}")
    print(f"Bereiche (keys): {len(area_agg)}")
    print(f"Blast-Muster:    {', '.join(blast_patterns)}")

    if args.export_tsv_dir:
        out_dir = Path(args.export_tsv_dir)
        out_dir.mkdir(parents=True, exist_ok=True)

        export_sheet_to_tsv(ws_summary, out_dir / "KW_Summary.tsv")
        export_sheet_to_tsv(ws_products, out_dir / "KW_Produkte.tsv")
        export_sheet_to_tsv(ws_areas, out_dir / "KW_Bereiche.tsv")
        print(f"TSV Export:      {out_dir}")

    print("=" * 90)


if __name__ == "__main__":
    main()
