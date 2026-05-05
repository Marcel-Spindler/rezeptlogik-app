#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
KW 19 ERWEITERTE ANALYSE: Detailliertes Tracking mit Positionen & Negationen
Trennt Zugänge (positiv) von Bewegungen/Rückgaben (negativ)
"""

from openpyxl import load_workbook
from collections import defaultdict
import csv

wb = load_workbook('Transaction_Log.xlsx')
ws = wb.active

# Lese Header
headers = []
for col in range(1, ws.max_column + 1):
    headers.append(ws.cell(row=1, column=col).value)

col_map = {h: i for i, h in enumerate(headers)}

print(f"{'='*120}")
print("KW 19 ERWEITERTE ANALYSE: PRODUKT-TRACKING MIT ZUGÄNGEN & BEWEGUNGEN")
print(f"{'='*120}\n")

# Sammle Daten mit getrennten positiven/negativen Mengen
products = defaultdict(lambda: {
    'bereich_positive': defaultdict(float),  # Zugänge
    'bereich_negative': defaultdict(float),  # Rückgaben/Bewegungen
    'bereich_netto': defaultdict(float),     # Netto = Positiv - Negativ
    'transaktionen_pos': 0,
    'transaktionen_neg': 0,
    'gesamt_pos': 0,
    'gesamt_neg': 0,
    'netto': 0,
    'daten': []
})

bereiche_set = set()
kw19_count = 0
positive_movements = 0
negative_movements = 0

print(f"Lese {ws.max_row - 1} Zeilen...")

for row_idx in range(2, ws.max_row + 1):
    if row_idx % 5000 == 0:
        print(f"  Verarbeitet: {row_idx}/{ws.max_row}...")
    
    week = ws.cell(row=row_idx, column=col_map['Week']+1).value
    if str(week) != '202619':
        continue
    
    kw19_count += 1
    
    item_desc = ws.cell(row=row_idx, column=col_map['Item Desc']+1).value
    if not item_desc:
        continue
    item_desc = str(item_desc).strip()
    
    tran_qty = ws.cell(row=row_idx, column=col_map['Tran Qty']+1).value
    try:
        qty = float(tran_qty) if tran_qty else 0
    except:
        qty = 0
    
    location_id = ws.cell(row=row_idx, column=col_map['Location Id']+1).value
    bereich = str(location_id).strip() if location_id else "UNBEKANNT"
    bereiche_set.add(bereich)
    
    tran_type = ws.cell(row=row_idx, column=col_map['Tran Type']+1).value
    description = ws.cell(row=row_idx, column=col_map['Description']+1).value
    user_zone = ws.cell(row=row_idx, column=col_map['User Zone']+1).value
    
    # Klassifiziere positive vs negative Bewegungen
    if qty >= 0:
        is_positive = True
        products[item_desc]['bereich_positive'][bereich] += qty
        products[item_desc]['transaktionen_pos'] += 1
        products[item_desc]['gesamt_pos'] += qty
        positive_movements += 1
    else:
        is_positive = False
        products[item_desc]['bereich_negative'][bereich] += abs(qty)
        products[item_desc]['transaktionen_neg'] += 1
        products[item_desc]['gesamt_neg'] += abs(qty)
        negative_movements += 1
    
    products[item_desc]['bereich_netto'][bereich] += qty
    products[item_desc]['netto'] += qty
    
    products[item_desc]['daten'].append({
        'typ': 'ZUGANG' if is_positive else 'BEWEGUNG',
        'bereich': bereich,
        'menge': abs(qty),
        'zone': user_zone,
        'tran_type': tran_type,
        'beschreibung': description
    })

print(f"✓ {kw19_count} Transaktionen KW 19")
print(f"  - Positive (Zugänge): {positive_movements}")
print(f"  - Negative (Bewegungen/Rückgaben): {negative_movements}\n")

# Sortiere nach Netto
sorted_products = sorted(products.items(), key=lambda x: abs(x[1]['netto']), reverse=True)

# Statistik
total_positiv = sum(p[1]['gesamt_pos'] for p in sorted_products)
total_negativ = sum(p[1]['gesamt_neg'] for p in sorted_products)
total_netto = sum(p[1]['netto'] for p in sorted_products)

print(f"{'='*120}")
print("GEWICHTS-STATISTIK KW 19:")
print(f"{'='*120}")
print(f"Gesamt ZUGÄNGE (positiv):     {total_positiv:>15.2f} kg")
print(f"Gesamt BEWEGUNGEN (negativ):  {total_negativ:>15.2f} kg")
print(f"NETTO-BESTAND:               {total_netto:>15.2f} kg")
print(f"Anzahl einzigartiger Produkte: {len(sorted_products)}\n")

print(f"{'='*120}")
print("TOP 30 PRODUKTE NACH NETTO-BESTAND")
print(f"{'='*120}\n")

print(f"{'Pos':<4} {'Produkt':<45} {'Zugänge':>12} {'Bewegungen':>12} {'NETTO':>12} {'Trans+':>8} {'Trans-':>8}")
print(f"{'-'*120}")

for i, (prod_name, data) in enumerate(sorted_products[:30], 1):
    print(f"{i:<4} {prod_name:<45} {data['gesamt_pos']:>12.2f} {data['gesamt_neg']:>12.2f} {data['netto']:>12.2f} {data['transaktionen_pos']:>8} {data['transaktionen_neg']:>8}")

# BEREICH-KATEGORISIERUNG
inbound_kw = ['inb', 'empf', 'warin', 'receiv', 'stage']
lager_kw = ['lager', 'storage', 'ambient', 'chilled', 'frozen', 'pick', 'slot', 'stgdr']
fulfil_kw = ['fab', 'fac', 'preb', 'pack', 'slip', 'vf-', 'plating', 'plh-', 'psh-', 'vegg']
prep_kw = ['preb', 'prep', 'plating', 'post', 'sleeving', 'deboxwip', 'platingwip']

print(f"\n\n{'='*120}")
print("DETAILLIERTE BEREICHS-ANALYSE (Nach Produkttyp)")
print(f"{'='*120}\n")

# Kategorisiere Bereiche
bereich_categories = {
    'Inbound': [],
    'Lager': [],
    'Prep/Plating': [],
    'Fulfillment': [],
    'Andere': []
}

for bereich in sorted(bereiche_set):
    bl = bereich.lower()
    if any(k in bl for k in inbound_kw):
        bereich_categories['Inbound'].append(bereich)
    elif any(k in bl for k in prep_kw):
        bereich_categories['Prep/Plating'].append(bereich)
    elif any(k in bl for k in fulfil_kw):
        bereich_categories['Fulfillment'].append(bereich)
    elif any(k in bl for k in lager_kw):
        bereich_categories['Lager'].append(bereich)
    else:
        bereich_categories['Andere'].append(bereich)

print("BEREICH-KATEGORIEN:")
for category, bereiche in bereich_categories.items():
    print(f"\n{category} ({len(bereiche)} Bereiche):")
    for b in bereiche[:5]:
        print(f"  - {b}")
    if len(bereiche) > 5:
        print(f"  ... und {len(bereiche)-5} weitere")

# Gewichtsfluss pro Kategorie
print(f"\n\n{'='*120}")
print("MATERIAL-FLOW NACH BEREICH-KATEGORIE")
print(f"{'='*120}\n")

print(f"{'Kategorie':<20} {'Zugänge':>15} {'Bewegungen':>15} {'NETTO':>15} {'Anteil':>8}")
print(f"{'-'*75}")

flow_by_category = {}
for category, bereiche in bereich_categories.items():
    total_pos = 0
    total_neg = 0
    total_net = 0
    
    for prod_name, data in sorted_products:
        for b in bereiche:
            total_pos += data['bereich_positive'].get(b, 0)
            total_neg += data['bereich_negative'].get(b, 0)
            total_net += data['bereich_netto'].get(b, 0)
    
    flow_by_category[category] = (total_pos, total_neg, total_net)
    
    if total_netto != 0:
        pct = (total_net / total_netto) * 100
    else:
        pct = 0
    
    print(f"{category:<20} {total_pos:>15.2f} {total_neg:>15.2f} {total_net:>15.2f} {pct:>7.1f}%")

# TOP PRODUKTE PRO KATEGORIE
print(f"\n\n{'='*120}")
print("TOP 5 PRODUKTE PRO BEREICH-KATEGORIE")
print(f"{'='*120}\n")

for category, bereiche in bereich_categories.items():
    print(f"\n{category}:")
    print(f"{'-'*80}")
    
    category_products = defaultdict(float)
    for prod_name, data in sorted_products:
        for b in bereiche:
            category_products[prod_name] += data['bereich_netto'].get(b, 0)
    
    sorted_cat = sorted(category_products.items(), key=lambda x: abs(x[1]), reverse=True)[:5]
    
    for prod_name, netto in sorted_cat:
        if netto != 0:
            print(f"  {prod_name:<60} {netto:>12.2f} kg")

# Exportiere zu erweiterte CSV
output_file = 'KW19_Analyse_Erweitert_Zugaenge_Bewegungen.csv'
print(f"\n\nExportiere zu {output_file}...")

with open(output_file, 'w', newline='', encoding='utf-8') as f:
    writer = csv.writer(f, delimiter=';')
    
    # Header
    all_bereiche = sorted(bereiche_set)
    header = ['Produkt', 'Gesamt_Zugänge_kg', 'Gesamt_Bewegungen_kg', 'NETTO_kg', 'Trans_Zugänge', 'Trans_Bewegungen']
    header += ['Zugänge_' + b for b in all_bereiche]
    header += ['Bewegungen_' + b for b in all_bereiche]
    header += ['Netto_' + b for b in all_bereiche]
    writer.writerow(header)
    
    # Daten
    for prod_name, data in sorted_products:
        row = [
            prod_name,
            f"{data['gesamt_pos']:.2f}",
            f"{data['gesamt_neg']:.2f}",
            f"{data['netto']:.2f}",
            str(data['transaktionen_pos']),
            str(data['transaktionen_neg'])
        ]
        
        for b in all_bereiche:
            row.append(f"{data['bereich_positive'].get(b, 0):.2f}")
        for b in all_bereiche:
            row.append(f"{data['bereich_negative'].get(b, 0):.2f}")
        for b in all_bereiche:
            row.append(f"{data['bereich_netto'].get(b, 0):.2f}")
        
        writer.writerow(row)

print(f"✓ Exportiert: {output_file}\n")

print(f"{'='*120}")
print("ANALYSE ABGESCHLOSSEN")
print(f"{'='*120}")
