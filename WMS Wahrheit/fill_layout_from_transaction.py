#!/usr/bin/env python
# -*- coding: utf-8 -*-

"""
Fuellt ein bestehendes Layout-Workbook (W20, W21, ...) mit Ist-Werten aus Transaction Logs.

Aktuell fokussiert auf Blast-Chiller-Ist (Spalte J) in den Wxx-Sheets:
- Keine Ueberschreibung bestehender Werte.
- KW wird aus Sheetname gelesen (z. B. W20 -> 202620).
- Produkt-Matching ueber Spalte C (Sub-Rezept/Ingredient Name).
- Quelle: aggregierte Mengen aus Blast-/PostB-Bereichen.
"""

from __future__ import annotations

import argparse
import re
from collections import defaultdict
from pathlib import Path
from typing import Dict, List, Tuple

from openpyxl import load_workbook


def to_float(value) -> float:
    if value is None or value == "":
        return 0.0
    try:
        return float(value)
    except Exception:
        return 0.0


def normalize_text(value: str) -> str:
    text = (value or "").strip().lower()
    text = re.sub(r"\s+", " ", text)
    return text


def normalize_week(value) -> str:
    if value is None:
        return ""
    s = str(value).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s


def category_for_area(area: str) -> str:
    a = normalize_text(area)
    prep_kw = ("preb", "prep", "plating", "post", "sleeving", "deboxwip", "platingwip")
    fulfil_kw = ("fab", "fac", "preb", "pack", "slip", "vf-", "plh-", "psh-", "vegg")
    lager_kw = ("lager", "storage", "ambient", "chilled", "frozen", "pick", "slot", "stgdr")

    if any(k in a for k in prep_kw):
        return "Prep/Plating"
    if any(k in a for k in fulfil_kw):
        return "Fulfillment"
    if any(k in a for k in lager_kw):
        return "Lager"
    return "Andere"


def aggregate_blast_by_week_and_product(transaction_log: Path, blast_patterns: List[str]) -> Dict[str, Dict[str, float]]:
    wb = load_workbook(transaction_log)
    ws = wb.active

    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    col_map = {str(h): i + 1 for i, h in enumerate(headers) if h is not None}

    required = ["Week", "Item Desc", "Tran Qty", "Location Id"]
    missing = [c for c in required if c not in col_map]
    if missing:
        raise ValueError(f"Pflichtspalten fehlen im Transaction Log: {missing}")

    patterns = [p.lower().strip() for p in blast_patterns if p and p.strip()]
    data = defaultdict(lambda: defaultdict(float))

    for row in range(2, ws.max_row + 1):
        week = normalize_week(ws.cell(row=row, column=col_map["Week"]).value)
        if not week:
            continue

        product_raw = ws.cell(row=row, column=col_map["Item Desc"]).value
        if product_raw is None:
            continue
        product = normalize_text(str(product_raw))
        if not product:
            continue

        qty = abs(to_float(ws.cell(row=row, column=col_map["Tran Qty"]).value))
        location_raw = ws.cell(row=row, column=col_map["Location Id"]).value
        location = normalize_text(str(location_raw) if location_raw else "")

        if any(p in location for p in patterns):
            data[week][product] += qty

    return data


def aggregate_department_weights(transaction_log: Path) -> Dict[str, Dict[Tuple[str, str], Dict[str, float]]]:
    wb = load_workbook(transaction_log)
    ws = wb.active

    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    col_map = {str(h): i + 1 for i, h in enumerate(headers) if h is not None}

    required = ["Week", "Tran Qty", "Location Id"]
    missing = [c for c in required if c not in col_map]
    if missing:
        raise ValueError(f"Pflichtspalten fehlen im Transaction Log: {missing}")

    out = defaultdict(lambda: defaultdict(lambda: {"pos": 0.0, "neg": 0.0, "net": 0.0}))

    for row in range(2, ws.max_row + 1):
        week = normalize_week(ws.cell(row=row, column=col_map["Week"]).value)
        if not week:
            continue

        qty = to_float(ws.cell(row=row, column=col_map["Tran Qty"]).value)
        area_raw = ws.cell(row=row, column=col_map["Location Id"]).value
        area = str(area_raw).strip() if area_raw else "UNBEKANNT"
        abteilung = category_for_area(area)

        key = (abteilung, area)
        if qty >= 0:
            out[week][key]["pos"] += qty
        else:
            out[week][key]["neg"] += abs(qty)
        out[week][key]["net"] += qty

    return out


def kw_sheet_to_weekcode(sheet_name: str, year: int) -> str | None:
    m = re.match(r"^W(\d{1,2})$", sheet_name.strip(), re.IGNORECASE)
    if not m:
        return None
    kw = int(m.group(1))
    return f"{year}{kw:02d}"


def find_value_for_name(product_name: str, week_map: Dict[str, float]) -> float:
    key = normalize_text(product_name)
    if not key:
        return 0.0

    if key in week_map:
        return week_map[key]

    # Fallback: enthaelt / enthalten
    for candidate, value in week_map.items():
        if key in candidate or candidate in key:
            return value

    return 0.0


def fill_layout(layout_file: Path, tx_file: Path, out_file: Path, year: int, blast_patterns: List[str]) -> Dict[str, int]:
    blast_data = aggregate_blast_by_week_and_product(tx_file, blast_patterns)
    dept_data = aggregate_department_weights(tx_file)
    wb = load_workbook(layout_file)

    stats = {
        "sheets": 0,
        "rows_scanned": 0,
        "filled": 0,
        "skipped_existing": 0,
        "not_found": 0,
    }

    for ws in wb.worksheets:
        week_code = kw_sheet_to_weekcode(ws.title, year)
        if not week_code:
            continue

        stats["sheets"] += 1
        week_map = blast_data.get(week_code, {})

        # Layout laut Datei: C=Sub-Rezept Name, D=WO, J=Blast Chiller Actual
        for r in range(5, ws.max_row + 1):
            recipe_part = ws.cell(r, 3).value
            wo = ws.cell(r, 4).value

            if recipe_part is None or str(recipe_part).strip() == "":
                continue
            if wo is None or str(wo).strip() == "":
                continue

            stats["rows_scanned"] += 1

            actual_cell = ws.cell(r, 10)  # J
            if actual_cell.value not in (None, ""):
                stats["skipped_existing"] += 1
                continue

            value = find_value_for_name(str(recipe_part), week_map)
            if value > 0:
                actual_cell.value = round(value, 2)
                stats["filled"] += 1
            else:
                stats["not_found"] += 1

    # Extra Reiter mit Gewichten pro Abteilung/Bereich je KW.
    sheet_name = "Abteilungsgewichte"
    if sheet_name in wb.sheetnames:
        ws_dept = wb[sheet_name]
        ws_dept.delete_rows(1, ws_dept.max_row)
    else:
        ws_dept = wb.create_sheet(sheet_name)

    headers = [
        "Week",
        "Abteilung",
        "Bereich",
        "Zugaenge_kg",
        "Bewegungen_kg",
        "Differenz_kg",
        "Netto_kg",
    ]
    for c, h in enumerate(headers, start=1):
        ws_dept.cell(1, c, h)

    row_no = 2
    for week in sorted(dept_data.keys()):
        entries = dept_data[week]
        sorted_keys = sorted(entries.keys(), key=lambda k: (k[0], -abs(entries[k]["net"]), k[1]))
        for abteilung, bereich in sorted_keys:
            vals = entries[(abteilung, bereich)]
            ws_dept.cell(row_no, 1, week)
            ws_dept.cell(row_no, 2, abteilung)
            ws_dept.cell(row_no, 3, bereich)
            ws_dept.cell(row_no, 4, round(vals["pos"], 2))
            ws_dept.cell(row_no, 5, round(vals["neg"], 2))
            ws_dept.cell(row_no, 6, round(vals["pos"] - vals["neg"], 2))
            ws_dept.cell(row_no, 7, round(vals["net"], 2))
            row_no += 1

    wb.save(out_file)
    return stats


def main():
    parser = argparse.ArgumentParser(description="Fuelle Layout-Workbook aus Transaction Log")
    parser.add_argument("--layout", required=True, help="Pfad zum Layout-Workbook")
    parser.add_argument("--transaction-log", required=True, help="Pfad zur Transaction_Log.xlsx")
    parser.add_argument("--output", default="", help="Ausgabedatei (optional)")
    parser.add_argument("--year", type=int, default=2026, help="Jahr fuer Wxx -> YYYYWW")
    parser.add_argument(
        "--blast-patterns",
        default="postb-01,blast,chill",
        help="Kommagetrennte Bereichsmuster fuer Blast-Chiller-Mengen",
    )

    args = parser.parse_args()
    layout = Path(args.layout)
    tx = Path(args.transaction_log)

    if not layout.exists():
        raise FileNotFoundError(f"Layout nicht gefunden: {layout}")
    if not tx.exists():
        raise FileNotFoundError(f"Transaction Log nicht gefunden: {tx}")

    out = Path(args.output) if args.output else layout.with_name(layout.stem + "_filled.xlsx")
    patterns = [p.strip() for p in args.blast_patterns.split(",") if p.strip()]

    stats = fill_layout(layout, tx, out, args.year, patterns)

    print("=" * 90)
    print("LAYOUT FUELLEN ABGESCHLOSSEN")
    print("=" * 90)
    print(f"Layout:         {layout}")
    print(f"TransactionLog: {tx}")
    print(f"Output:         {out}")
    print(f"Wxx-Sheets:     {stats['sheets']}")
    print(f"Zeilen geprueft:{stats['rows_scanned']}")
    print(f"Neu gefuellt:   {stats['filled']}")
    print(f"Schon befuellt: {stats['skipped_existing']}")
    print(f"Nicht gefunden: {stats['not_found']}")
    print(f"Patterns:       {', '.join(patterns)}")
    print("=" * 90)


if __name__ == "__main__":
    main()
