import openpyxl
from openpyxl.utils import get_column_letter

draft_path = "DRAFT VolumeBuffer Overview F_ .xlsx"
wb = openpyxl.load_workbook(draft_path, data_only=True)
ws = wb.active

# Print ALL non-empty rows with column references
print("=== ALL non-empty rows in draft ===")
for r in range(1, ws.max_row+1):
    row = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column+1)]
    if any(v not in (None,"") for v in row):
        # compact print: only columns A..N (1..14) since right side seems metadata
        compact = [(get_column_letter(c+1), v) for c,v in enumerate(row[:14]) if v not in (None,"")]
        meta = [(get_column_letter(c+1), v) for c,v in enumerate(row[14:], start=14) if v not in (None,"")]
        print(f"R{r}: {compact}  | meta: {meta}")
