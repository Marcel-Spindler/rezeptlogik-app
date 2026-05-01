import openpyxl
from collections import defaultdict

inv_path = "Current_Inventory (2).xlsx"
wb = openpyxl.load_workbook(inv_path, data_only=True)
ws = wb.active

# header
hdr = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column+1)]
idx = {h: i for i, h in enumerate(hdr)}
print("idx:", idx)

# search items
keywords = [
    "spinach", "stamppot",
    "green bean",
    "honey-garlic", "carrot coin", "carrot",
    "creamy dijon", "dijon",
    "salmon",
    "broccoli",
    "caper", "kapern",
    "cauliflower", "parmesan",
    "piccata",
]

# group inventory rows by Item Description, accumulate Actual Quantity by UoM
groups = defaultdict(lambda: defaultdict(list))  # desc -> uom -> [qty,...]
descs_set = set()
for r in range(2, ws.max_row+1):
    desc = ws.cell(row=r, column=idx['Item Description']+1).value
    uom = ws.cell(row=r, column=idx['Unit of Measure']+1).value
    qty = ws.cell(row=r, column=idx['Actual Quantity']+1).value
    if desc is None: continue
    try:
        q = float(qty)
    except (TypeError, ValueError):
        continue
    groups[desc][uom].append(q)
    descs_set.add(desc.lower())

print("\n=== Matches per keyword ===")
for kw in keywords:
    matches = [d for d in groups if kw.lower() in d.lower()]
    if matches:
        print(f"\n--- '{kw}' ---")
        for d in matches:
            for uom, lst in groups[d].items():
                print(f"  [{uom}] {d}  -> n={len(lst)}, sum={sum(lst):.2f}, vals={lst[:8]}")
