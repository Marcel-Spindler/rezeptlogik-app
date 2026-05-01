import openpyxl
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook(r"c:\Rezeptlogik\F_EU - 2026 Ramp Up Planning V2.0 (1).xlsx", data_only=True)

print("=== SHEETS ===")
for s in wb.sheetnames:
    ws = wb[s]
    print(f"  {s!r}  dims={ws.max_row}x{ws.max_column}")

target_sheets = [s for s in wb.sheetnames if "meal" in s.lower() or "selection" in s.lower() or "benl" in s.lower() or "dksde" in s.lower() or "dkse" in s.lower() or s.upper() == "DE"]
if not target_sheets:
    target_sheets = wb.sheetnames

for s in target_sheets:
    ws = wb[s]
    print(f"\n\n========== SHEET: {s} ({ws.max_row}x{ws.max_column}) ==========")
    rows_to_show = min(ws.max_row, 30)
    cols_to_show = min(ws.max_column, 30)
    for r in range(1, rows_to_show + 1):
        line = []
        for c in range(1, cols_to_show + 1):
            v = ws.cell(r, c).value
            if v is None:
                line.append("")
            else:
                s_v = str(v).replace("\t", " ").replace("\n", " ")
                if len(s_v) > 30:
                    s_v = s_v[:27] + "..."
                line.append(s_v)
        print(f"R{r:>3} | " + " | ".join(line))
